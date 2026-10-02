"""Gerador de documentos Word (.docx) — cartas, contratos com edição automática."""

import os
import io
from typing import Dict, Any, Optional

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "uploads", "documentos_word")
os.makedirs(OUTPUT_DIR, exist_ok=True)

def gerar_documento_word(titulo: str, conteudo: str, dados: Dict[str, Any]) -> Optional[str]:
    """
    Gera documento Word a partir de template com placeholders.
    Substitui {{CHAVE}} pelos valores em dados.
    """
    if not HAS_DOCX:
        return None

    # Substituir placeholders
    conteudo_preenchido = conteudo
    for chave, valor in dados.items():
        placeholder = "{" + "{" + chave + "}" + "}"
        conteudo_preenchido = conteudo_preenchido.replace(placeholder, str(valor))

    # Criar documento
    doc = Document()

    # Cabeçalho
    header = doc.add_heading(titulo, 0)
    header.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    header_run = header.runs[0]
    header_run.font.size = Pt(16)
    header_run.font.bold = True
    header_run.font.color.rgb = RGBColor(0, 102, 204)

    # Linha horizontal (simulated)
    doc.add_paragraph("_" * 80).alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # Conteúdo principal
    for paragrafo_texto in conteudo_preenchido.split('\n'):
        if paragrafo_texto.strip():
            p = doc.add_paragraph(paragrafo_texto)
            p_format = p.paragraph_format
            p_format.line_spacing = 1.5
            p_format.space_after = Pt(6)

            # Formatar runs com negrito se detectar títulos
            if ':' in paragrafo_texto and len(paragrafo_texto) < 100:
                for run in p.runs:
                    run.font.bold = True

            for run in p.runs:
                run.font.size = Pt(11)
        else:
            doc.add_paragraph()  # Espaço em branco

    # Rodapé
    doc.add_paragraph()
    rodape = doc.add_paragraph("Documento gerado automaticamente por on.imob")
    rodape.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    rodape_run = rodape.runs[0]
    rodape_run.font.size = Pt(9)
    rodape_run.font.italic = True
    rodape_run.font.color.rgb = RGBColor(153, 153, 153)

    # Salvar
    filename = f"{titulo.replace(' ', '_').lower()}.docx"
    filepath = os.path.join(OUTPUT_DIR, filename)

    try:
        doc.save(filepath)
        return filepath
    except Exception as e:
        print(f"Erro ao salvar Word: {e}")
        return None

def gerar_em_memoria(titulo: str, conteudo: str, dados: Dict[str, Any]) -> Optional[bytes]:
    """
    Gera documento Word em memória (para download direto).
    Retorna bytes do arquivo .docx
    """
    if not HAS_DOCX:
        return None

    # Substituir placeholders
    conteudo_preenchido = conteudo
    for chave, valor in dados.items():
        placeholder = "{" + "{" + chave + "}" + "}"
        conteudo_preenchido = conteudo_preenchido.replace(placeholder, str(valor))

    # Criar documento
    doc = Document()

    # Cabeçalho
    header = doc.add_heading(titulo, 0)
    header.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    header_run = header.runs[0]
    header_run.font.size = Pt(16)
    header_run.font.bold = True
    header_run.font.color.rgb = RGBColor(0, 102, 204)

    doc.add_paragraph("_" * 80).alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # Conteúdo
    for paragrafo_texto in conteudo_preenchido.split('\n'):
        if paragrafo_texto.strip():
            p = doc.add_paragraph(paragrafo_texto)
            p_format = p.paragraph_format
            p_format.line_spacing = 1.5
            p_format.space_after = Pt(6)

            if ':' in paragrafo_texto and len(paragrafo_texto) < 100:
                for run in p.runs:
                    run.font.bold = True

            for run in p.runs:
                run.font.size = Pt(11)
        else:
            doc.add_paragraph()

    # Rodapé
    doc.add_paragraph()
    rodape = doc.add_paragraph("Documento gerado automaticamente por on.imob")
    rodape.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    rodape_run = rodape.runs[0]
    rodape_run.font.size = Pt(9)
    rodape_run.font.italic = True
    rodape_run.font.color.rgb = RGBColor(153, 153, 153)

    # Retornar em memória
    buffer = io.BytesIO()
    try:
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
    except Exception as e:
        print(f"Erro ao gerar Word em memória: {e}")
        return None
